import csv, time, datetime, requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── Cấu hình ────────────────────────────────────────────────────
TZ_VN = datetime.timezone(datetime.timedelta(hours=7))

# Endpoint xác nhận chạy được cho MỌI platform (FB/IG/Zalo, khả năng cao cả TikTok/Shopee)
# miễn là dùng đúng "Public API access token" riêng của từng page.
BASE_URL = "https://pages.fm/api/public_api/v1"

PAGES_CSV   = "pages_config_full.csv"    # cột: platform,channel,page_id,access_token
OUTPUT_XLSX = "chat_stats_output.xlsx"
DELAY_BETWEEN_CALLS = 0.8

SESSION = requests.Session()

def load_pages(path: str) -> list[dict]:
    pages = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            pages.append({
                "platform":     row["platform"].strip().lower(),
                "channel":      row["channel"].strip(),
                "page_id":      row["page_id"].strip(),
                "access_token": row.get("access_token", "").strip(),
            })
    return pages

def date_to_ts(date_str: str) -> tuple[int, int]:
    dt = datetime.datetime.strptime(date_str, "%d/%m/%Y").replace(tzinfo=TZ_VN)
    since = int(dt.timestamp())
    until = int(dt.replace(hour=23, minute=59, second=59).timestamp())
    return since, until

def fetch_chat_stats(page_id: str, token: str, date_str: str, retries=3) -> dict | None:
    """Trả về dict {total, qc, organic} cho 1 page (mọi platform) trong 1 ngày."""
    since, until = date_to_ts(date_str)
    url = f"{BASE_URL}/pages/{page_id}/statistics/ads"

    res = None
    for attempt in range(retries):
        try:
            res = SESSION.get(url, params={
                "access_token":   token,
                "since":          since,
                "until":          until,
                "type_statistic": "day",
                "type":           "by_time",
            }, timeout=30)
            break
        except requests.RequestException as e:
            print(f"   ⚠️  Lần {attempt+1}/{retries} lỗi mạng: {e}")
            if attempt < retries - 1:
                time.sleep(5 * (attempt + 1))
                continue
            return None

    if res is None:
        return None

    if res.status_code == 429:
        print("   ⏳ Bị rate-limit (429), nghỉ 10s rồi thử lại...")
        time.sleep(10)
        try:
            res = SESSION.get(url, params={
                "access_token":   token,
                "since":          since,
                "until":          until,
                "type_statistic": "day",
                "type":           "by_time",
            }, timeout=30)
        except requests.RequestException as e:
            print(f"   ❌ Lỗi mạng sau retry: {e}")
            return None

    if res.status_code != 200:
        print(f"   ❌ Pancake {res.status_code}: {res.text[:200]}")
        return None

    data = res.json()
    if data.get("success") is False:
        print(f"   ❌ Pancake báo lỗi: {data.get('message')}")
        return None

    page_stats = data.get("page_statistic", [])
    total_inbox = sum(r.get("new_inbox_count", 0) for r in page_stats)

    ads_data = data.get("data", [])
    if isinstance(ads_data, dict):
        ads_data = ads_data.get("data", [])
    total_ads_click = sum(r.get("sum_click", 0) for r in ads_data)

    organic = max(total_inbox - total_ads_click, 0)

    return {"total": total_inbox, "qc": total_ads_click, "organic": organic}

# ── Main ──────────────────────────────────────────────────────
def run_export():
    t0 = time.time()
    yesterday = (datetime.datetime.now(TZ_VN) - datetime.timedelta(days=1)).strftime("%d/%m/%Y")

    pages = load_pages(PAGES_CSV)
    print(f"\n🚀 Xuất dữ liệu ngày: {yesterday} | {len(pages)} kênh\n")

    rows = []
    ok = fail = 0

    for i, p in enumerate(pages, start=1):
        channel, platform, page_id, token = p["channel"], p["platform"], p["page_id"], p["access_token"]
        print(f"[{i}/{len(pages)}] 🔄 {channel} ({platform})")

        if not token:
            print("   ⚠️  Thiếu access_token trong CSV — bỏ qua.")
            rows.append({"channel": channel, "platform": platform, "page_id": page_id,
                          "total": "", "qc": "", "organic": "", "status": "thiếu access_token"})
            fail += 1
        else:
            stats = fetch_chat_stats(page_id, token, yesterday)
            if stats is None:
                rows.append({"channel": channel, "platform": platform, "page_id": page_id,
                              "total": "", "qc": "", "organic": "", "status": "lỗi fetch"})
                fail += 1
            else:
                print(f"   Tổng: {stats['total']} | QC: {stats['qc']} | Organic: {stats['organic']}")
                rows.append({"channel": channel, "platform": platform, "page_id": page_id,
                              "total": stats["total"], "qc": stats["qc"], "organic": stats["organic"],
                              "status": "ok"})
                ok += 1

        if i < len(pages):
            time.sleep(DELAY_BETWEEN_CALLS)

    write_xlsx(rows, yesterday, OUTPUT_XLSX)

    print(f"\n📄 Đã xuất kết quả ra: {OUTPUT_XLSX}")
    print(f"🏁 Xong trong {time.time()-t0:.1f}s | ✅ {ok} thành công | ❌ {fail} thất bại")


def write_xlsx(rows: list[dict], date_str: str, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Chat Stats"

    headers = ["Kênh", "Nền tảng", "Page ID", "Tổng hội thoại", "Hội thoại QC", "Hội thoại Organic", "Trạng thái"]
    ws.append(headers)

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    status_colors = {"ok": "C6EFCE", "thiếu access_token": "FFEB9C", "lỗi fetch": "FFC7CE"}

    for row in rows:
        ws.append([row["channel"], row["platform"], row["page_id"],
                   row["total"], row["qc"], row["organic"], row["status"]])
        r = ws.max_row
        fill_color = status_colors.get(row["status"], "FFFFFF")
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col_idx)
            cell.font = Font(name="Arial")
            cell.fill = fill

    for col_idx, header in enumerate(headers, start=1):
        max_len = len(header)
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    ws.freeze_panes = "A2"
    note_row = ws.max_row + 2
    ws.cell(row=note_row, column=1, value=f"Dữ liệu ngày: {date_str}").font = Font(name="Arial", italic=True, size=9)

    wb.save(path)

if __name__ == "__main__":
    run_export()